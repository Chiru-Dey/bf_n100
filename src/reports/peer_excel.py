"""Peer comparison Excel report generation with colour-coded percentiles."""

from __future__ import annotations

import logging
import sqlite3
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src.analytics.peer import METRICS_TO_RANK, load_peer_groups
from src.analytics.ratios import DB_PATH

logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_PATH = PROJECT_ROOT / "output" / "peer_comparison.xlsx"

GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
YELLOW_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
GOLD_FILL = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

METRIC_HEADERS = {
    "return_on_equity_pct": "ROE (%)",
    "return_on_capital_employed_pct": "ROCE (%)",
    "net_profit_margin_pct": "NPM (%)",
    "debt_to_equity": "D/E",
    "free_cash_flow_cr": "FCF (Cr)",
    "pat_cagr_5yr": "PAT CAGR 5Y (%)",
    "revenue_cagr_5yr": "Rev CAGR 5Y (%)",
    "eps_cagr_5yr": "EPS CAGR 5Y (%)",
    "interest_coverage": "ICR",
    "asset_turnover": "Asset Turnover",
}


def fetch_peer_data(
    db_path: Path = DB_PATH,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Fetch latest ratios, company names, and peer percentiles from the database."""
    with sqlite3.connect(db_path) as conn:
        ratios = pd.read_sql_query("SELECT * FROM financial_ratios", conn)
        companies = pd.read_sql_query("SELECT id, company_name FROM companies", conn)
        percentiles = pd.read_sql_query("SELECT * FROM peer_percentiles", conn)

    march = ratios[ratios["year"].astype(str).str.endswith("-03")]
    latest = march.sort_values("year").groupby("company_id").tail(1)
    missing = set(ratios["company_id"]) - set(latest["company_id"])
    if missing:
        fallback = (
            ratios[ratios["company_id"].isin(missing)]
            .sort_values("year")
            .groupby("company_id")
            .tail(1)
        )
        latest = pd.concat([latest, fallback])

    return latest, companies, percentiles


def build_group_sheet(
    group_name: str,
    members: list[str],
    benchmark: str | None,
    ratios: pd.DataFrame,
    companies: pd.DataFrame,
    percentiles: pd.DataFrame,
) -> tuple[pd.DataFrame, str | None]:
    """Build the DataFrame for a single peer group sheet including median row."""
    group_ratios = ratios[ratios["company_id"].isin(members)].copy()
    group_ratios = group_ratios.merge(
        companies.rename(columns={"id": "company_id"}), on="company_id", how="left"
    )

    group_pcts = percentiles[
        (percentiles["peer_group_name"] == group_name)
        & (percentiles["metric"].isin(METRICS_TO_RANK))
    ].pivot(index="company_id", columns="metric", values="percentile_rank")
    group_pcts.columns = [f"{c}_pct" for c in group_pcts.columns]
    group_pcts = group_pcts.reset_index()

    merged = group_ratios.merge(group_pcts, on="company_id", how="left")

    cols = ["company_id", "company_name"]
    for metric in METRICS_TO_RANK:
        if metric in merged.columns:
            cols.append(metric)
        pct_col = f"{metric}_pct"
        if pct_col in merged.columns:
            cols.append(pct_col)
    merged = merged[cols]

    median_row = {"company_id": "MEDIAN", "company_name": ""}
    for metric in METRICS_TO_RANK:
        if metric in merged.columns:
            median_row[metric] = merged[metric].median()

    return pd.concat([merged, pd.DataFrame([median_row])], ignore_index=True), benchmark


def export_peer_comparison_excel(
    path: Path = OUTPUT_PATH, db_path: Path = DB_PATH
) -> Path:
    """Generate the peer_comparison.xlsx file with 11 sheets."""
    peer_groups = load_peer_groups()
    ratios, companies, percentiles = fetch_peer_data(db_path)

    wb = Workbook()
    if wb.active:
        wb.remove(wb.active)

    for group_name, group_df in peer_groups.groupby("peer_group_name"):
        members = group_df["company_id"].tolist()
        benchmarks = group_df[group_df["is_benchmark"] == 1]["company_id"].tolist()
        benchmark = benchmarks[0] if benchmarks else None

        sheet_df, benchmark_id = build_group_sheet(
            group_name, members, benchmark, ratios, companies, percentiles
        )

        ws = wb.create_sheet(title=group_name[:31])

        headers = ["Company ID", "Company Name"]
        for metric in METRICS_TO_RANK:
            if metric in sheet_df.columns:
                headers.append(METRIC_HEADERS.get(metric, metric))
            pct_col = f"{metric}_pct"
            if pct_col in sheet_df.columns:
                headers.append(f"{METRIC_HEADERS.get(metric, metric)} Pct")

        for c_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=c_idx, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        for r_idx, row in enumerate(sheet_df.itertuples(index=False), 2):
            is_benchmark_row = row.company_id == benchmark_id
            is_median_row = row.company_id == "MEDIAN"

            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx)
                cell.value = None if pd.isna(value) else value

                if is_benchmark_row:
                    cell.fill = GOLD_FILL
                    cell.font = Font(bold=True)
                elif is_median_row:
                    cell.font = Font(italic=True, bold=True)

                col_name = headers[c_idx - 1]
                if col_name.endswith(" Pct") and not pd.isna(value):
                    try:
                        pct_val = float(value)
                        if pct_val >= 0.75:
                            cell.fill = GREEN_FILL
                        elif pct_val >= 0.25:
                            cell.fill = YELLOW_FILL
                        else:
                            cell.fill = RED_FILL
                    except (ValueError, TypeError):
                        pass

                if isinstance(value, float) and not pd.isna(value):
                    if "Pct" in col_name or col_name in [
                        "D/E",
                        "Asset Turnover",
                        "ICR",
                    ]:
                        cell.number_format = "0.00"
                    else:
                        cell.number_format = "#,##0.00"

        for c_idx in range(1, len(headers) + 1):
            ws.column_dimensions[get_column_letter(c_idx)].width = 16

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    logger.info("Wrote peer comparison report to %s", path)
    return path


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    export_peer_comparison_excel()
