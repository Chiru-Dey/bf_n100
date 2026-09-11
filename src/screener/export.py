"""Screener Excel export with threshold-based colour coding."""

from __future__ import annotations

import logging
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

from src.screener.engine import (
    FINANCIALS_SECTOR,
    METRIC_MAP,
    load_screener_config,
    run_preset,
)

logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parents[2]
OUTPUT_PATH = PROJECT_ROOT / "output" / "screener_output.xlsx"

GREEN_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
RED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")

EXPORT_COLUMNS = [
    "company_id",
    "broad_sector",
    "composite_quality_score",
    "return_on_equity_pct",
    "debt_to_equity",
    "free_cash_flow_cr",
    "revenue_cagr_5yr",
    "pat_cagr_5yr",
    "operating_profit_margin_pct",
    "pe_ratio",
    "pb_ratio",
    "dividend_yield_pct",
    "interest_coverage",
    "market_cap_crore",
    "net_profit",
    "eps_cagr_5yr",
    "asset_turnover",
    "sales",
    "revenue_cagr_3yr",
    "dividend_payout_ratio_pct",
]


def _check_threshold(
    value: float, key: str, threshold: float, sector: str
) -> bool | None:
    """Return True if value passes, False if fails, None if N/A or Financials D/E skip."""
    if pd.isna(value):
        return None
    if key not in METRIC_MAP:
        return None
    if key == "de_max" and sector == FINANCIALS_SECTOR:
        return None

    _, op = METRIC_MAP[key]
    if op == ">=":
        return float(value) >= threshold
    if op == "<=":
        return float(value) <= threshold
    if op == "==":
        return float(value) == threshold
    return None


def export_screener_excel(path: Path = OUTPUT_PATH) -> Path:
    """Generate the multi-sheet screener output Excel file with colour-coded cells."""
    config = load_screener_config()
    presets = config.get("presets", {})

    wb = Workbook()
    if wb.active:
        wb.remove(wb.active)

    for preset_name, filters in presets.items():
        df = run_preset(preset_name)
        if df.empty:
            continue

        cols = [c for c in EXPORT_COLUMNS if c in df.columns]
        sheet_df = df[cols].copy()

        ws = wb.create_sheet(title=preset_name[:31])

        for r_idx, row in enumerate(
            dataframe_to_rows(sheet_df, index=False, header=True), 1
        ):
            for c_idx, value in enumerate(row, 1):
                cell = ws.cell(row=r_idx, column=c_idx, value=value)

                if r_idx > 1:
                    col_name = cols[c_idx - 1]
                    sector = (
                        sheet_df.iloc[r_idx - 2]["broad_sector"]
                        if "broad_sector" in sheet_df.columns
                        else ""
                    )

                    for filter_key, threshold in filters.items():
                        if (
                            filter_key in METRIC_MAP
                            and METRIC_MAP[filter_key][0] == col_name
                        ):
                            passes = _check_threshold(
                                value, filter_key, threshold, sector
                            )
                            if passes is True:
                                cell.fill = GREEN_FILL
                            elif passes is False:
                                cell.fill = RED_FILL
                            break

        for col_idx, col_name in enumerate(cols, 1):
            max_len = max(len(str(col_name)), 12)
            ws.column_dimensions[ws.cell(row=1, column=col_idx).column_letter].width = (
                max_len + 2
            )

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    logger.info("Wrote screener output to %s", path)
    return path


if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)
    export_screener_excel()
