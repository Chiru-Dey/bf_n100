"""Unit tests for the peer comparison Excel report."""

from pathlib import Path

from openpyxl import load_workbook

from src.reports.peer_excel import export_peer_comparison_excel


def test_export_peer_comparison_excel_creates_file(tmp_path: Path) -> None:
    out_path = tmp_path / "test_peer.xlsx"
    export_peer_comparison_excel(out_path)
    assert out_path.exists()
    assert out_path.stat().st_size > 1000

    wb = load_workbook(out_path)
    assert len(wb.sheetnames) == 11
    assert "IT Services" in wb.sheetnames or "Private Banks" in wb.sheetnames