"""Unit tests for the screener Excel export module."""

from pathlib import Path

from openpyxl import load_workbook

from src.screener.export import export_screener_excel


def test_export_screener_excel_creates_file(tmp_path: Path) -> None:
    out_path = tmp_path / "test_screener.xlsx"
    export_screener_excel(out_path)
    assert out_path.exists()
    assert out_path.stat().st_size > 1000

    wb = load_workbook(out_path)
    assert len(wb.sheetnames) == 6
    assert "quality_compounder" in wb.sheetnames
    assert "value_pick" in wb.sheetnames
